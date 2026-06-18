"""文档解析异步任务测试."""

from typing import Any

import pytest
from sqlalchemy.orm import Session

from app.database import SessionLocal
from app.models.document import Document
from app.models.financial_report import FinancialReport
from app.models.tenant import Tenant
from app.models.user import User
from app.parser.simple_parser import SimpleDocumentParser
from app.security import get_password_hash
from app.tasks.document_tasks import parse_document_task


class FakeStorageClient:
    """测试用对象存储客户端."""

    def __init__(self, stored: dict[str, bytes] | None = None) -> None:
        """初始化."""
        self.stored = stored or {}

    def upload_bytes(
        self,
        key: str,
        data: bytes,
        _content_type: str = "application/octet-stream",
        _metadata: dict[str, Any] | None = None,
    ) -> str:
        """模拟上传."""
        self.stored[key] = data
        return f"http://fake-minio/financial-agent/{key}"

    def download_bytes(self, key: str) -> bytes:
        """模拟下载."""
        return self.stored[key]


def _create_document(db: Session, tenant: Tenant, user: User, filename: str) -> Document:
    """辅助函数：创建并提交测试文档."""
    doc = Document(
        tenant_id=tenant.id,
        created_by=user.id,
        filename=filename,
        storage_key=f"docs/{filename}",
        status="pending",
    )
    db.add(doc)
    db.commit()
    db.refresh(doc)
    return doc


def test_parse_document_task_success(monkeypatch: pytest.MonkeyPatch) -> None:
    """测试文档解析任务成功更新状态与结果.

    注意：Celery 任务使用独立的 SessionLocal，为避免 SQLite 事务隔离导致任务看不到
    未提交数据，本测试在独立会话中创建租户、用户与文档并真实提交。
    """
    db = SessionLocal()
    try:
        tenant = Tenant(name="Task Test Tenant", code="task-test")
        db.add(tenant)
        db.commit()
        db.refresh(tenant)

        user = User(
            tenant_id=tenant.id,
            username="tasktester",
            email="task@example.com",
            hashed_password=get_password_hash("testpass"),
            role="admin",
            is_active="Y",
        )
        db.add(user)
        db.commit()
        db.refresh(user)

        storage_key = "docs/profit_q2_2025.txt"
        fake_storage = FakeStorageClient({storage_key: b""})
        monkeypatch.setattr(
            "app.tasks.document_tasks.get_storage_client",
            lambda: fake_storage,
        )

        doc = _create_document(db, tenant, user, "profit_q2_2025.txt")
        doc.storage_key = storage_key
        db.commit()

        result = parse_document_task.delay(doc.id).get()

        assert result["status"] == "success"
        assert result["document_id"] == doc.id

        db.refresh(doc)
        assert doc.status == "success"
        assert doc.confidence == 0.3
        assert doc.parse_result is not None
        assert doc.parse_result["detected_year"] == 2025
        assert doc.parse_result["detected_period"] == "Q2"
        assert doc.error_message is None
    finally:
        db.close()


def test_parse_document_task_not_found() -> None:
    """测试对不存在文档执行任务返回失败但不重试."""
    result = parse_document_task.delay("non-existent-id").get()

    assert result["status"] == "failed"
    assert result["retry"] is False


def test_simple_parser_extracts_metadata() -> None:
    """测试简单解析器从文件名提取元数据."""
    doc = Document(
        tenant_id="t1",
        created_by="u1",
        filename="balance_2024_h1.xlsx",
        storage_key="docs/balance_2024_h1.xlsx",
        status="pending",
    )
    parser = SimpleDocumentParser(doc)

    result = parser.parse(b"", doc.filename)

    assert result["extension"] == "xlsx"
    assert result["detected_year"] == 2024
    assert result["detected_period"] == "H1"
    assert result["confidence"] == 0.3


def _make_excel_bytes(rows: list[list[Any]]) -> bytes:
    """构造一个简单 xlsx 文件字节流."""
    from io import BytesIO

    from openpyxl import Workbook

    wb = Workbook()
    ws = wb.active
    if ws is None:
        raise RuntimeError("No active worksheet")
    for row in rows:
        ws.append(row)

    buffer = BytesIO()
    wb.save(buffer)
    return buffer.getvalue()


def test_parse_excel_document_imports_financial_data(
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    """测试 Excel 文档解析后导入 financial_reports."""
    db = SessionLocal()
    try:
        tenant = Tenant(name="Excel Import Tenant", code="excel-import")
        db.add(tenant)
        db.commit()
        db.refresh(tenant)

        user = User(
            tenant_id=tenant.id,
            username="exceltester",
            email="excel@example.com",
            hashed_password=get_password_hash("testpass"),
            role="admin",
            is_active="Y",
        )
        db.add(user)
        db.commit()
        db.refresh(user)

        storage_key = "docs/profit_2025_q2.xlsx"
        excel_content = _make_excel_bytes(
            [
                ["year", "period", "revenue", "operating_cost", "net_profit"],
                [2025, "Q2", 10_000_000, 6_000_000, 2_500_000],
            ]
        )
        fake_storage = FakeStorageClient({storage_key: excel_content})
        monkeypatch.setattr(
            "app.tasks.document_tasks.get_storage_client",
            lambda: fake_storage,
        )

        doc = _create_document(db, tenant, user, "profit_2025_q2.xlsx")
        doc.storage_key = storage_key
        db.commit()

        result = parse_document_task.delay(doc.id).get()

        assert result["status"] == "success"

        db.refresh(doc)
        assert doc.status == "success"
        assert doc.parse_result is not None
        assert doc.parse_result["format"] == "excel"
        assert doc.parse_result["imported_count"] == 1

        report = (
            db.query(FinancialReport)
            .filter(
                FinancialReport.tenant_id == tenant.id,
                FinancialReport.year == 2025,
                FinancialReport.period == "Q2",
            )
            .first()
        )
        assert report is not None
        assert report.revenue == 10_000_000.0
        assert report.operating_cost == 6_000_000.0
        assert report.net_profit == 2_500_000.0
    finally:
        db.close()


def test_parse_csv_document_imports_financial_data(
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    """测试 CSV 文档解析后导入 financial_reports."""
    db = SessionLocal()
    try:
        tenant = Tenant(name="CSV Import Tenant", code="csv-import")
        db.add(tenant)
        db.commit()
        db.refresh(tenant)

        user = User(
            tenant_id=tenant.id,
            username="csvtester",
            email="csv@example.com",
            hashed_password=get_password_hash("testpass"),
            role="admin",
            is_active="Y",
        )
        db.add(user)
        db.commit()
        db.refresh(user)

        storage_key = "docs/profit_2025_q2.csv"
        csv_content = (
            b"year,period,revenue,operating_cost,net_profit\n2025,Q2,10000000,6000000,2500000\n"
        )
        fake_storage = FakeStorageClient({storage_key: csv_content})
        monkeypatch.setattr(
            "app.tasks.document_tasks.get_storage_client",
            lambda: fake_storage,
        )

        doc = _create_document(db, tenant, user, "profit_2025_q2.csv")
        # 确保 storage_key 与 fake storage 一致
        doc.storage_key = storage_key
        db.commit()

        result = parse_document_task.delay(doc.id).get()

        assert result["status"] == "success"

        db.refresh(doc)
        assert doc.status == "success"
        assert doc.parse_result is not None
        assert doc.parse_result["format"] == "csv"
        assert doc.parse_result["imported_count"] == 1

        report = (
            db.query(FinancialReport)
            .filter(
                FinancialReport.tenant_id == tenant.id,
                FinancialReport.year == 2025,
                FinancialReport.period == "Q2",
            )
            .first()
        )
        assert report is not None
        assert report.revenue == 10_000_000.0
        assert report.operating_cost == 6_000_000.0
        assert report.net_profit == 2_500_000.0
    finally:
        db.close()
